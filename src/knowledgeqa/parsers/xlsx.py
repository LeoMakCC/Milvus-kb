from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from knowledgeqa.ingestion.models import ParsedDocument, ParsedSection


class XlsxParser:
    extensions = {".xlsx", ".xlsm"}

    def parse(self, path: Path) -> ParsedDocument:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[ParsedSection] = []
        try:
            for worksheet in workbook.worksheets:
                lines: list[str] = []
                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        str(value).strip()
                        for value in row
                        if value is not None and str(value).strip()
                    ]
                    if values:
                        lines.append(" | ".join(values))
                if lines:
                    section_name = f"工作表: {worksheet.title}"
                    sections.append(
                        ParsedSection(
                            text=f"{section_name}\n" + "\n".join(lines),
                            section=section_name,
                        )
                    )
        finally:
            workbook.close()
        text = "\n\n".join(section.text for section in sections)
        return ParsedDocument(
            title=path.stem,
            text=text,
            doc_type="xlsx",
            sections=sections,
        )
