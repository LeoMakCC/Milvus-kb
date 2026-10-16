from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from knowledgeqa.ingestion.models import ParsedDocument, ParsedSection


class PptxParser:
    extensions = {".pptx"}

    def parse(self, path: Path) -> ParsedDocument:
        presentation = Presentation(path)
        sections: list[ParsedSection] = []
        for slide_no, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        texts.append(text)
            if texts:
                sections.append(
                    ParsedSection(
                        text="\n".join(texts),
                        section=f"幻灯片 {slide_no}",
                        page_no=slide_no,
                    )
                )
        text = "\n\n".join(section.text for section in sections)
        return ParsedDocument(
            title=path.stem,
            text=text,
            doc_type="pptx",
            sections=sections,
        )
