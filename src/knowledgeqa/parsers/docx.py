from __future__ import annotations

from pathlib import Path

from docx import Document

from knowledgeqa.ingestion.models import ParsedDocument, ParsedSection


class DocxParser:
    extensions = {".docx"}

    def parse(self, path: Path) -> ParsedDocument:
        document = Document(path)
        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        text = "\n".join(paragraphs)
        title = document.core_properties.title or path.stem
        return ParsedDocument(
            title=title,
            text=text,
            doc_type="docx",
            sections=[ParsedSection(text=text)] if text else [],
        )
