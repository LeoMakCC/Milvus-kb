from pathlib import Path

import pymupdf
import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation

from knowledgeqa.parsers.base import OcrRequiredError
from knowledgeqa.parsers.registry import parse_document


def test_pdf_parser_extracts_pages_with_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Hello PDF")
    pdf.save(path)
    pdf.close()

    document = parse_document(path)

    assert document.doc_type == "pdf"
    assert "Hello PDF" in document.text
    assert document.sections[0].page_no == 1


def test_blank_pdf_requires_ocr(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()

    with pytest.raises(OcrRequiredError):
        parse_document(path)


def test_docx_parser_extracts_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = WordDocument()
    document.add_paragraph("Hello Word")
    document.save(path)

    parsed = parse_document(path)

    assert parsed.doc_type == "docx"
    assert "Hello Word" in parsed.text


def test_xlsx_parser_extracts_sheet_rows(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合同"
    sheet.append(["编号", "金额"])
    sheet.append(["HT-1", 100])
    workbook.save(path)

    parsed = parse_document(path)

    assert parsed.doc_type == "xlsx"
    assert "工作表: 合同" in parsed.text
    assert "编号 | 金额" in parsed.text
    assert "HT-1 | 100" in parsed.text


def test_pptx_parser_extracts_slide_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Hello Slide"
    presentation.save(path)

    parsed = parse_document(path)

    assert parsed.doc_type == "pptx"
    assert "Hello Slide" in parsed.text
    assert parsed.sections[0].page_no == 1
